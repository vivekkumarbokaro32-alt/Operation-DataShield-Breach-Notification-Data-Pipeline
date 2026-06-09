"""
Operation DataShield — Phase 7: Client Report Generation
Goal: Produce a fully formatted Excel report for MedCore's
      legal, compliance, and executive teams.

Sheets:
  1. Executive Summary   — headline numbers + SLA countdown
  2. Volume by State     — record counts per US state
  3. Volume by Source    — which vendor contributed what
  4. Exception Detail    — the NO_CONTACT records
  5. Pipeline Audit      — what happened at each phase
"""

import pandas as pd
from sqlalchemy import create_engine, text
from openpyxl import load_workbook
from openpyxl.styles import (PatternFill, Font, Alignment,
                              Border, Side)
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
import logging
from datetime import datetime, date
import os

# ─────────────────────────────────────────────────────────────
# SETUP
# ─────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s | %(levelname)s | %(message)s',
    handlers=[
        logging.FileHandler('phase7_report.log'),
        logging.StreamHandler()
    ]
)
log = logging.getLogger('phase7')

DB_URL = "postgresql+psycopg2://postgres:Vivek%40827009@localhost:5432/Datasheild"
engine = create_engine(DB_URL)

# Project dates
PROJECT_START = date(2026, 6, 8)
SLA_DAYS      = 60
SLA_DEADLINE  = date(2026, 8, 7)
TODAY         = date.today()
DAYS_ELAPSED  = (TODAY - PROJECT_START).days
DAYS_REMAINING = SLA_DAYS - DAYS_ELAPSED

# Output path
OUTPUT_DIR  = 'output'
REPORT_PATH = os.path.join(
    OUTPUT_DIR,
    f"DataShield_Client_Report_{TODAY.strftime('%Y-%m-%d')}.xlsx"
)

# ─────────────────────────────────────────────────────────────
# COLOR SCHEME — consistent across all sheets
# ─────────────────────────────────────────────────────────────
C_DARK_BLUE  = '1F3864'   # headers
C_MID_BLUE   = '2E75B6'   # subheaders
C_LIGHT_BLUE = 'D6E4F0'   # alternate rows
C_GREEN      = '70AD47'   # good/ok indicators
C_AMBER      = 'FFC000'   # warning indicators
C_RED        = 'FF0000'   # critical indicators
C_WHITE      = 'FFFFFF'
C_LIGHT_GREY = 'F2F2F2'

def header_style(cell, dark=True):
    cell.fill = PatternFill('solid',
                 fgColor=C_DARK_BLUE if dark else C_MID_BLUE)
    cell.font = Font(color=C_WHITE, bold=True, size=11)
    cell.alignment = Alignment(horizontal='center',
                                vertical='center',
                                wrap_text=True)

def auto_width(ws, min_w=12, max_w=45):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = \
            max(min_w, min(max_len + 4, max_w))

def thin_border():
    s = Side(style='thin', color='CCCCCC')
    return Border(left=s, right=s, top=s, bottom=s)

# ─────────────────────────────────────────────────────────────
# DATA LOADING
# ─────────────────────────────────────────────────────────────
def load_data() -> dict:
    """Load all required data from PostgreSQL."""
    log.info("Loading data from database...")
    with engine.connect() as conn:
        df_ready  = pd.read_sql(
            text("SELECT * FROM processed.notification_ready"), conn)
        df_exc    = pd.read_sql(
            text("SELECT * FROM processed.exceptions"), conn)
        df_norm   = pd.read_sql(
            text("SELECT record_source, COUNT(*) as count "
                 "FROM staging.normalized GROUP BY record_source"), conn)
        df_deduped = pd.read_sql(
            text("SELECT COUNT(*) as count "
                 "FROM staging.deduped"), conn)

    log.info(f"  notification_ready: {len(df_ready):,}")
    log.info(f"  exceptions:         {len(df_exc):,}")
    return {
        'ready':         df_ready,
        'exceptions':    df_exc,
        'normalized':    df_norm,
        'deduped_count': df_deduped['count'].iloc[0]
    }

# ─────────────────────────────────────────────────────────────
# SHEET 1: Executive Summary
# ─────────────────────────────────────────────────────────────
def build_executive_summary(ws, data: dict):
    """Build the Executive Summary sheet."""
    log.info("  Building Executive Summary...")
    df = data['ready']
    df_exc = data['exceptions']

    # SLA status color
    if DAYS_REMAINING > 20:
        sla_color = C_GREEN
        sla_status = 'ON TRACK'
    elif DAYS_REMAINING > 7:
        sla_color = C_AMBER
        sla_status = 'MONITOR'
    else:
        sla_color = C_RED
        sla_status = 'AT RISK'

    # Channel counts
    channel_counts = df['channel'].value_counts().to_dict()

    rows = [
        # SLA Banner
        ['OPERATION DATASHIELD — NOTIFICATION STATUS REPORT', ''],
        ['Client: MedCore Health Systems', ''],
        ['Prepared by: EPIQ Data Operations', ''],
        [f'Report Date: {TODAY.strftime("%B %d, %Y")}', ''],
        ['', ''],

        # SLA Section
        ['SLA STATUS', ''],
        ['Project Start Date',     PROJECT_START.strftime('%B %d, %Y')],
        ['Legal Deadline (HIPAA)', SLA_DEADLINE.strftime('%B %d, %Y')],
        ['Days Elapsed',           DAYS_ELAPSED],
        ['Days Remaining',         DAYS_REMAINING],
        ['SLA Status',             sla_status],
        ['', ''],

        # Population Section
        ['AFFECTED POPULATION SUMMARY', ''],
        ['Raw Records Received (3 vendors)',   21500],
        ['Duplicate Records Removed',          21500 - data['deduped_count']],
        ['Duplication Rate',
         f"{round((21500 - data['deduped_count'])/21500*100, 1)}%"],
        ['Unique Individuals Identified',      data['deduped_count']],
        ['', ''],

        # Notification Section
        ['NOTIFICATION BREAKDOWN', ''],
        ['Total Notification Ready',
         len(df)],
        ['Email + Mail (BOTH)',
         channel_counts.get('BOTH', 0)],
        ['Mail Only',
         channel_counts.get('MAIL_ONLY', 0)],
        ['Email Only',
         channel_counts.get('EMAIL_ONLY', 0)],
        ['No Contact Info (Escalated)',
         len(df_exc)],
        ['', ''],

        # Output Files Section
        ['OUTPUT FILES GENERATED', ''],
        ['Print/Mail File (PrintCo)',
         channel_counts.get('BOTH', 0) +
         channel_counts.get('MAIL_ONLY', 0)],
        ['Email File (SendGrid)',
         channel_counts.get('BOTH', 0) +
         channel_counts.get('EMAIL_ONLY', 0)],
        ['Call Center File',        len(df[df['phone_valid'] == True])],
        ['Credit Monitoring File',  len(df[df['dob_flag'] == 'OK'])],
    ]

    for r_idx, row in enumerate(rows, 1):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)

            # Title row
            if r_idx == 1:
                cell.font = Font(bold=True, size=14,
                                 color=C_DARK_BLUE)
                cell.alignment = Alignment(horizontal='left')

            # Section headers
            elif str(val).isupper() and c_idx == 1 and val != '':
                cell.fill = PatternFill('solid', fgColor=C_DARK_BLUE)
                cell.font = Font(bold=True, color=C_WHITE, size=11)

            # SLA status cell
            elif val == sla_status:
                cell.fill = PatternFill('solid', fgColor=sla_color)
                cell.font = Font(bold=True, color=C_WHITE, size=11)
                cell.alignment = Alignment(horizontal='center')

            # Value column
            elif c_idx == 2 and val != '':
                cell.font = Font(size=11)
                cell.alignment = Alignment(horizontal='right')
                if isinstance(val, int):
                    cell.number_format = '#,##0'

            # Label column
            elif c_idx == 1:
                cell.font = Font(size=11)

    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 20
    ws.freeze_panes = 'A2'


# ─────────────────────────────────────────────────────────────
# SHEET 2: Volume by State
# ─────────────────────────────────────────────────────────────
def build_state_volume(ws, data: dict):
    """Build the Volume by State sheet."""
    log.info("  Building Volume by State...")
    df = data['ready']

    state_counts = df.groupby('state_code').agg(
        total_records=('record_id', 'count'),
        email_count=('email_valid', 'sum'),
        mail_count=('address_valid', 'sum')
    ).reset_index()
    state_counts.columns = ['State', 'Total Records',
                             'Email Count', 'Mail Count']
    state_counts = state_counts.sort_values(
        'Total Records', ascending=False)
    state_counts['% of Total'] = (
        state_counts['Total Records'] /
        state_counts['Total Records'].sum() * 100
    ).round(1).astype(str) + '%'

    # Write header
    headers = ['State', 'Total Records',
               'Email Count', 'Mail Count', '% of Total']
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        header_style(cell)

    # Write data
    for r_idx, row in enumerate(
            state_counts.itertuples(index=False), 2):
        for c_idx, val in enumerate(
                [row.State, row._1, row._2, row._3, row._4], 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border()
            if r_idx % 2 == 0:
                cell.fill = PatternFill('solid',
                                        fgColor=C_LIGHT_BLUE)
            if isinstance(val, (int, float)) and c_idx != 1:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right')

    # Totals row
    total_row = len(state_counts) + 2
    ws.cell(row=total_row, column=1, value='TOTAL').font = \
        Font(bold=True)
    total_cell = ws.cell(
        row=total_row, column=2,
        value=state_counts['Total Records'].sum())
    total_cell.font = Font(bold=True)
    total_cell.number_format = '#,##0'

    auto_width(ws)
    ws.freeze_panes = 'A2'


# ─────────────────────────────────────────────────────────────
# SHEET 3: Volume by Vendor Source
# ─────────────────────────────────────────────────────────────
def build_source_volume(ws, data: dict):
    """Build the Volume by Vendor Source sheet."""
    log.info("  Building Volume by Source...")
    df = data['ready']
    df_norm = data['normalized']

    # Raw counts from normalized table
    raw_counts = df_norm.set_index('record_source')['count'].to_dict()

    # Post-dedup counts from notification_ready
    deduped_counts = df['record_source'].value_counts().to_dict()

    headers = ['Vendor Source', 'Raw Records',
               'After Dedup', 'Records Removed', 'Reduction %']
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        header_style(cell)

    vendors = ['vendor_a', 'vendor_b', 'vendor_c']
    for r_idx, vendor in enumerate(vendors, 2):
        raw    = raw_counts.get(vendor, 0)
        deduped = deduped_counts.get(vendor, 0)
        removed = raw - deduped
        pct     = round(removed / raw * 100, 1) if raw > 0 else 0

        vals = [vendor.replace('_', ' ').title(),
                raw, deduped, removed, f"{pct}%"]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border()
            if r_idx % 2 == 0:
                cell.fill = PatternFill('solid',
                                        fgColor=C_LIGHT_BLUE)
            if isinstance(val, int):
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right')

    auto_width(ws)
    ws.freeze_panes = 'A2'


# ─────────────────────────────────────────────────────────────
# SHEET 4: Exception Detail
# ─────────────────────────────────────────────────────────────
def build_exception_detail(ws, data: dict):
    """Build the Exception Detail sheet."""
    log.info("  Building Exception Detail...")
    df_exc = data['exceptions']

    if df_exc.empty:
        ws.cell(row=1, column=1,
                value='No exceptions found.')
        return

    # Select and rename relevant columns
    cols = ['record_id', 'first_name', 'last_name',
            'email', 'phone_e164', 'address_1',
            'city', 'state_code', 'zip5',
            'email_flag', 'phone_flag',
            'address_flag', 'channel', 'record_source']
    df_out = df_exc[[c for c in cols if c in df_exc.columns]]

    headers = [c.replace('_', ' ').title()
               for c in df_out.columns]
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        header_style(cell)

    for r_idx, row in enumerate(df_out.itertuples(index=False), 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border()
            cell.fill = PatternFill('solid', fgColor='FFF2CC')

    # Add note at top
    ws.insert_rows(1)
    note = ws.cell(
        row=1, column=1,
        value=(f"ACTION REQUIRED: {len(df_exc)} records have no "
               f"valid contact information. "
               f"Please provide alternative contact details "
               f"or confirm these individuals cannot be reached."))
    note.font = Font(bold=True, color=C_RED, size=11)
    note.fill = PatternFill('solid', fgColor='FFE0E0')
    ws.merge_cells(f'A1:{get_column_letter(len(headers))}1')

    auto_width(ws)
    ws.freeze_panes = 'A3'


# ─────────────────────────────────────────────────────────────
# SHEET 5: Pipeline Audit
# ─────────────────────────────────────────────────────────────
def build_pipeline_audit(ws, data: dict):
    """Build the Pipeline Audit sheet."""
    log.info("  Building Pipeline Audit...")

    deduped = data['deduped_count']
    ready   = len(data['ready'])
    exc     = len(data['exceptions'])

    audit_rows = [
        ['Phase', 'Description',
         'Records In', 'Records Out', 'Notes'],
        ['Phase 1', 'Raw Ingestion',
         21500, 21500,
         '3 vendor files: A(8,000) B(7,500) C(6,000)'],
        ['Phase 2', 'Data Profiling',
         21500, 21500,
         '4 DOB formats, 5 phone formats detected'],
        ['Phase 3', 'Normalization',
         21500, 21500,
         'Names, state, zip, phone, DOB standardized'],
        ['Phase 4', 'Deduplication',
         21500, deduped,
         f'{21500-deduped:,} duplicates removed (3 passes)'],
        ['Phase 5', 'Validation',
         deduped, ready,
         f'{exc} NO_CONTACT exceptions escalated'],
        ['Phase 6', 'Vendor Output',
         ready, ready,
         '4 vendor files + manifest generated'],
    ]

    for r_idx, row in enumerate(audit_rows, 1):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border()

            if r_idx == 1:
                header_style(cell)
            else:
                if r_idx % 2 == 0:
                    cell.fill = PatternFill('solid',
                                            fgColor=C_LIGHT_BLUE)
                if isinstance(val, int):
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal='right')
                # Highlight Phase 4 reduction
                if r_idx == 5 and c_idx == 4:
                    cell.font = Font(bold=True, color=C_MID_BLUE)

    auto_width(ws)
    ws.freeze_panes = 'A2'


# ─────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────
if __name__ == '__main__':
    log.info("Operation DataShield — Phase 7: Report Generation")
    log.info(f"Run timestamp: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    log.info(f"SLA: {DAYS_REMAINING} days remaining")

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # ── Load all data ─────────────────────────────────────────
    data = load_data()

    # ── Build Excel with pandas first ────────────────────────
    log.info("Building Excel report...")

    # Create placeholder DataFrames for each sheet
    with pd.ExcelWriter(REPORT_PATH, engine='openpyxl') as writer:
        pd.DataFrame({'': []}).to_excel(
            writer, sheet_name='Executive Summary', index=False)
        pd.DataFrame({'': []}).to_excel(
            writer, sheet_name='Volume by State', index=False)
        pd.DataFrame({'': []}).to_excel(
            writer, sheet_name='Volume by Source', index=False)
        pd.DataFrame({'': []}).to_excel(
            writer, sheet_name='Exception Detail', index=False)
        pd.DataFrame({'': []}).to_excel(
            writer, sheet_name='Pipeline Audit', index=False)

    # ── Open with openpyxl and build each sheet ───────────────
    wb = load_workbook(REPORT_PATH)

    build_executive_summary(wb['Executive Summary'], data)
    build_state_volume(wb['Volume by State'], data)
    build_source_volume(wb['Volume by Source'], data)
    build_exception_detail(wb['Exception Detail'], data)
    build_pipeline_audit(wb['Pipeline Audit'], data)

    # ── Tab colors ────────────────────────────────────────────
    wb['Executive Summary'].sheet_properties.tabColor  = C_DARK_BLUE
    wb['Volume by State'].sheet_properties.tabColor    = C_MID_BLUE
    wb['Volume by Source'].sheet_properties.tabColor   = C_MID_BLUE
    wb['Exception Detail'].sheet_properties.tabColor   = 'FF0000'
    wb['Pipeline Audit'].sheet_properties.tabColor     = C_GREEN

    # ── Set Executive Summary as first active sheet ───────────
    wb.active = wb['Executive Summary']

    wb.save(REPORT_PATH)

    # ── Final summary ─────────────────────────────────────────
    file_size = os.path.getsize(REPORT_PATH)
    log.info("")
    log.info("=" * 60)
    log.info("PHASE 7 COMPLETE — CLIENT REPORT GENERATED")
    log.info("=" * 60)
    log.info(f"  Report:        {REPORT_PATH}")
    log.info(f"  File size:     {file_size/1024:.1f} KB")
    log.info(f"  Sheets:        5")
    log.info(f"  SLA status:    {DAYS_REMAINING} days remaining")
    log.info("")
    log.info("  FINAL PROJECT SUMMARY:")
    log.info(f"  Raw records received:    21,500")
    log.info(f"  Unique individuals:      {data['deduped_count']:,}")
    log.info(f"  Notification ready:      {len(data['ready']):,}")
    log.info(f"  Exceptions escalated:    {len(data['exceptions']):,}")
    log.info(f"  Vendor files generated:  4")
    log.info("=" * 60)
    log.info("Operation DataShield pipeline complete.")